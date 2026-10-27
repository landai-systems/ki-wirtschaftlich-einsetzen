"""Grenzfälle für „Mein Vorhaben“: setzt Eingaben, lässt LibreOffice neu rechnen und prüft gegen die Kontrollrechnung.
Aufruf: python grenzfaelle.py <rechenmodell.xlsx> <pfad zu recalc.py>"""
import shutil, subprocess, sys, tempfile, os
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from kontrollrechnung import pruefe

FAELLE = [  # Spalte: {Zeile: Wert}, erwartete Einschätzung (Anfang), erwarteter Hinweis (Teiltext)
    ({6: 0}, "Keine Vorgänge", "Vorgänge fehlen"),
    ({8: 7}, "Rechnet sich nicht", "nicht schneller"),
    ({8: 5.9}, "Rechnet sich nicht", "keine"),
    ({22: 0}, "Kalkulatorisch lohnend. Geld spart", "keine"),
    ({6: 100}, "Laufend positiv, aber", "keine"),
    ({9: 0.6, 7: 300, 11: 10, 12: 2}, None, "sehr hoch"),
]
src, recalc = sys.argv[1], sys.argv[2]
fehler = 0
for start in range(0, len(FAELLE), 3):
    tmp = os.path.join(tempfile.mkdtemp(), "g.xlsx"); shutil.copy(src, tmp)
    wb = openpyxl.load_workbook(tmp); ws = wb["Mein Vorhaben"]
    gruppe = FAELLE[start:start + 3]
    for col, (eing, _, _) in zip("BCD", gruppe):
        for r, v in eing.items():
            ws[f"{col}{r}"] = v
    wb.save(tmp)
    subprocess.run([sys.executable, recalc, tmp, "90"], check=True, capture_output=True)
    fehler += pruefe(tmp)
    ws = openpyxl.load_workbook(tmp, data_only=True)["Mein Vorhaben"]
    for col, (eing, einsch, hinweis) in zip("BCD", gruppe):
        e, h = ws[f"{col}45"].value, ws[f"{col}46"].value
        ok = (einsch is None or e.startswith(einsch)) and hinweis in h
        fehler += not ok
        print("OK    " if ok else "FEHLER", eing, "|", e, "|", h)
print("Fehler:", fehler); sys.exit(1 if fehler else 0)
