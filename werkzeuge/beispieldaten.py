"""Erzeugt die anonymisierten, konstruierten Beispieldaten des Unternehmerpakets."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
W = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MIN = [4.5, 7, 5.5, 6, 8.5, 5, 6.5, 4, 7.5, 6, 5.5, 6, 9, 4.5, 6, 5, 7, 6.5, 5.5, 6, 4, 8, 6, 5, 7, 6.5, 5.5, 6, 5, 5.5]
assert abs(sum(MIN) / len(MIN) - 6) < 1e-9, sum(MIN) / len(MIN)
b, n, f = Font(name="Arial", bold=True), Font(name="Arial", italic=True, color="555555", size=9), PatternFill("solid", fgColor="D9E1F2")
wb = Workbook(); ws = wb.active; ws.title = "Zeitmessung"
ws["A1"] = "Beispiel: Zeitmessung für 30 Kundenanfragen (konstruierte Werte, keine echten Unternehmensdaten)"; ws["A1"].font = b
ws["A2"] = "So messen: 30 gewöhnliche Vorgänge, Zeit vom Öffnen bis zum Abschluss, einfache und schwierige Fälle gemischt."; ws["A2"].font = n
for i, h in enumerate(["Nr.", "Vorgangsart", "Minuten"], 1):
    c = ws.cell(4, i, h); c.font = b; c.fill = f
for i, m in enumerate(MIN, 1):
    ws.cell(4 + i, 1, i); ws.cell(4 + i, 2, "Anfrage Haustür" if i % 3 else "Anfrage Fenster"); ws.cell(4 + i, 3, m)
e = 4 + len(MIN)
for k, (t, fml) in enumerate([("Anzahl", f"=COUNT(C5:C{e})"), ("Durchschnitt (in den Rechner: Minuten heute)", f"=AVERAGE(C5:C{e})"),
                              ("Median", f"=MEDIAN(C5:C{e})"), ("Kürzester / längster Vorgang", f'=MIN(C5:C{e})&" / "&MAX(C5:C{e})')]):
    ws.cell(e + 2 + k, 2, t).font = b; ws.cell(e + 2 + k, 3, fml)
ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 46; ws.column_dimensions["C"].width = 12
w2 = wb.create_sheet("Vorgänge zählen")
w2["A1"] = "Beispiel: Kundenanfragen eines Monats nach Kalenderwoche (konstruierte Werte)"; w2["A1"].font = b
for i, h in enumerate(["Woche", "Anfragen"], 1):
    c = w2.cell(3, i, h); c.font = b; c.fill = f
for i, (kw, a) in enumerate([("Woche 1", 142), ("Woche 2", 155), ("Woche 3", 149), ("Woche 4", 154)], 4):
    w2.cell(i, 1, kw); w2.cell(i, 2, a)
w2["A9"] = "Summe (in den Rechner: Vorgänge je Monat)"; w2["A9"].font = b; w2["B9"] = "=SUM(B4:B7)"
w2.column_dimensions["A"].width = 44
wb.save(os.path.join(W, "unternehmer", "Beispieldaten", "Beispiel_Ausgangsmessung.xlsx"))
