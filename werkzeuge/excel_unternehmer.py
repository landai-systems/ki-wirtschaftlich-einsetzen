"""Ergänzt das Rechenmodell des Buches um ein Startblatt und den Rechner „Mein Vorhaben“,
kennzeichnet Eingaben einheitlich und schützt Formeln (Blattschutz ohne Kennwort).
Aufruf: python excel_unternehmer.py <eingabe.xlsx> <ausgabe.xlsx>
Danach mit LibreOffice neu berechnen (recalc.py), damit die Ergebniswerte gespeichert sind."""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

F = "Arial"
f_title = Font(name=F, size=18, bold=True, color="1F3864")
f_h2 = Font(name=F, size=12, bold=True, color="1F3864")
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_note = Font(name=F, size=9, italic=True, color="555555")
f_link = Font(name=F, size=11, bold=True, color="0563C1", underline="single")
fill_in = PatternFill("solid", fgColor="DDEBF7")
fill_res = PatternFill("solid", fgColor="F2F2F2")
fill_head = PatternFill("solid", fgColor="D9E1F2")
fill_key = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
EUR = '#,##0.00 "€";-#,##0.00 "€";"0,00 €"'
EUR0 = '#,##0 "€";-#,##0 "€";"0 €"'
EUR4 = '#,##0.0000 "€"'
H1 = '#,##0.0" h";-#,##0.0" h";"0,0 h"'
INT = '#,##0'
DEC1 = '#,##0.0'
PCT = '0%'

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)

# 1. Division durch null im Buchmodell absichern (null Vorgänge)
w = wb["Wirtschaftlichkeit"]
for c in "BCD":
    w[f"{c}36"] = f"=IF({c}13>0,({c}30*Stundensatz+{c}34)/{c}13,\"-\")"

# 2. Rechner „Mein Vorhaben“
ws = wb.create_sheet("Mein Vorhaben", 0)
ws.sheet_view.showGridLines = False
for col, wd in zip("ABCDEF", (52, 15, 15, 15, 11, 62)):
    ws.column_dimensions[col].width = wd
ws["A1"] = "Mein Vorhaben: Rechnet sich die KI-Anwendung?"; ws["A1"].font = f_title
ws["A2"] = ("Blau hinterlegte Felder sind Eingaben: tragen Sie Ihre eigenen Werte ein. Grau hinterlegte Felder "
            "rechnet die Tabelle selbst. Alle Beträge netto. Vorbelegt ist das Beispiel aus Kapitel 5 des Buches "
            "(mittleres Volumen); „Vorsichtig“ halbiert die Zeitersparnis wie die Empfindlichkeitsrechnung im Buch.")
ws["A2"].font = f_note; ws["A2"].alignment = wrap; ws.merge_cells("A2:F2"); ws.row_dimensions[2].height = 40
ws["A3"] = "Rechenweg und Begriffe: Kapitel 5 und Anhang A des Buches. Ergebnisse sind Modellrechnungen, keine Zusage."
ws["A3"].font = f_note
hdr = ["Eingabe", "Vorsichtig", "Erwartet", "Ihre Variante", "Einheit", "Erläuterung"]
for i, h in enumerate(hdr):
    c = ws.cell(5, i + 1, h); c.font = f_bold; c.fill = fill_head; c.border = box
K = "Wirtschaftlichkeit!$C$10"
EING = [  # Zeile, Text, (vorsichtig, erwartet, variante), Einheit, Format, Erläuterung, Prüfung
    (6, "Vorgänge je Monat", (600, 600, 600), "Anzahl", INT, "Wie oft fällt der Vorgang im Monat an? Am besten zählen statt schätzen.", ("whole", 0, 1000000)),
    (7, "Minuten je Vorgang heute (ohne KI)", (6, 6, 6), "min", DEC1, "Gemessene Bearbeitungszeit heute, gemittelt über einfache und schwierige Fälle.", ("decimal", 0.1, 600)),
    (8, "Minuten je Vorgang mit KI (Prüfen und Freigeben)", (4, 2, 2), "min", DEC1, "Zeit, die ein Mensch mit KI noch braucht. Vorsichtig: halbe Zeitersparnis.", ("decimal", 0, 600)),
    (9, "Anteil der Vorgänge mit Nacharbeit", (0.08, 0.08, 0.08), "%", PCT, "Wie oft muss das KI-Ergebnis korrigiert werden? Im Pilotprojekt messen.", ("decimal", 0, 1)),
    (10, "Minuten je Nacharbeit", (5, 5, 5), "min", DEC1, "Zusätzliche Zeit, wenn korrigiert werden muss.", ("decimal", 0, 600)),
    (11, "Arbeitskosten je geleistete Stunde", (45, 45, 45), "€/h", EUR, "Bruttolohn plus Arbeitgeberanteile, geteilt durch tatsächlich geleistete Stunden (Blatt Arbeitskosten).", ("decimal", 1, 1000)),
    (12, "KI-Kosten je Vorgang", ("=" + K, "=" + K, "=" + K), "€", EUR4, "Vorbelegt aus dem Buchbeispiel (rund 1 Cent). Für eigene Werte: Blatt Vorgangsrechner oder Angebot des Dienstleisters.", ("decimal", 0, 100)),
    (13, "Feste Software- und Plattformkosten", (50, 50, 50), "€/Monat", EUR, "Monatliche Gebühren, z. B. für einen Automatisierungsdienst.", ("decimal", 0, 1000000)),
    (14, "Wartung und Kontrolle", (4, 4, 4), "h/Monat", DEC1, "Stichproben, Anpassungen, Störungen.", ("decimal", 0, 1000)),
    (15, "Neutest nach Modell- oder Anbieterwechsel", (8, 8, 8), "h/Jahr", DEC1, "KI-Modelle werden regelmäßig ersetzt; danach muss neu getestet werden.", ("decimal", 0, 1000)),
    (16, "Einrichtung durch Dienstleister", (50, 50, 50), "h", DEC1, "Stunden laut Angebot. Richtet ein eigener Mitarbeiter ein: 0 und die Stunden unten bei den internen Stunden eintragen.", ("decimal", 0, 10000)),
    (17, "Stundensatz Dienstleister", (100, 100, 100), "€/h", EUR, "Netto-Stundensatz laut Angebot.", ("decimal", 0, 1000)),
    (18, "Datenschutz und Recht, einmalig", (600, 600, 600), "€", EUR0, "Prüfung Auftragsverarbeitung, Verzeichnis, Information, gegebenenfalls Betriebsrat.", ("decimal", 0, 1000000)),
    (19, "IT-Sicherheit, einmalig", (400, 400, 400), "€", EUR0, "Zugänge, Rechte, Protokollierung.", ("decimal", 0, 1000000)),
    (20, "Schulung", (12, 12, 12), "h gesamt", DEC1, "Summe der Schulungsstunden aller Beteiligten.", ("decimal", 0, 10000)),
    (21, "Interne Stunden für Auswahl, Testfälle und Pilotprojekt", (16, 16, 16), "h gesamt", DEC1, "Bei eigener Einrichtung zusätzlich deren Stunden.", ("decimal", 0, 10000)),
    (22, "Anteil der frei werdenden Zeit, der tatsächlich Ausgaben vermeidet", (0.3, 0.3, 0.3), "%", PCT, "Nur wenn Überstunden, Aushilfen oder Neueinstellungen wegfallen, spart frei werdende Zeit Geld. 0 % = reiner Kapazitätsgewinn.", ("decimal", 0, 1)),
    (23, "Zusätzlicher Deckungsbeitrag je Monat (optional)", (0, 0, 0), "€/Monat", EUR, "Nur Deckungsbeitrag, nicht Umsatz: z. B. bisher vergessene Leistungen, die bezahlt werden, abzüglich Material. Im Zweifel 0.", ("decimal", 0, 1000000)),
    (24, "Betrachtungszeitraum", (24, 24, 24), "Monate", INT, "KI-Modelle und Preise ändern sich schnell; mehr als 36 Monate sind unsicher.", ("whole", 1, 120)),
]
for r, txt, vals, unit, fmt, note, (typ, lo, hi) in EING:
    a = ws.cell(r, 1, txt); a.font = f_norm; a.alignment = Alignment(wrap_text=True, vertical="center")
    for j, v in enumerate(vals):
        c = ws.cell(r, 2 + j, v); c.font = f_in; c.fill = fill_in; c.number_format = fmt; c.border = box
        c.protection = openpyxl.styles.Protection(locked=False)
    ws.cell(r, 5, unit).font = f_norm
    n = ws.cell(r, 6, note); n.font = f_note; n.alignment = wrap
    dv = DataValidation(type=typ, operator="between", formula1=str(lo), formula2=str(hi), allow_blank=False,
                        showErrorMessage=True, errorTitle="Unplausible Eingabe",
                        error=f"Bitte einen Wert zwischen {lo:g} und {hi:g} eingeben ({unit}). Prozentwerte als z. B. 30 %.")
    ws.add_data_validation(dv); dv.add(f"B{r}:D{r}")
    ws.row_dimensions[r].height = 26

c = ws.cell(26, 1, "Ergebnis"); c.font = f_bold; c.fill = fill_head
for j in range(2, 7):
    ws.cell(26, j).fill = fill_head
ERG = [  # Zeile, Text, Formel mit X als Spaltenplatzhalter, Format, Erläuterung, Schlüssel
    (27, "Arbeitszeit heute", "=X6*X7/60", H1, "Vorgänge × Minuten heute.", False),
    (28, "Arbeitszeit mit KI inkl. Nacharbeit und Wartung", "=(X6*X8+X6*X9*X10)/60+X14+X15/12", H1, "Prüfen, Nacharbeit, Wartung und anteiliger Neutest.", False),
    (29, "Frei werdende Arbeitszeit (Kapazitätsgewinn)", "=X27-X28", H1, "Zeit, die für andere Aufgaben frei wird. Das ist noch kein Geld.", False),
    (30, "Kapazitätsgewinn, mit Arbeitskosten bewertet", "=X29*X11", EUR, "Kalkulatorischer Wert der frei werdenden Zeit.", False),
    (31, "Laufende Kosten (KI und Software)", "=X6*X12+X13", EUR, "Zahlungswirksame Kosten je Monat.", False),
    (32, "Nettoeffekt kalkulatorisch je Monat", "=X30-X31", EUR, "Bewertete Zeitersparnis minus laufende Kosten.", True),
    (33, "Davon tatsächlich vermiedene Auszahlungen", "=X30*X22", EUR, "Nur der Anteil aus Zeile 22.", False),
    (34, "Nettoeffekt zahlungswirksam je Monat", "=X33+X23-X31", EUR, "Was am Monatsende tatsächlich auf dem Konto bleibt (inkl. Deckungsbeitrag aus Zeile 23).", True),
    (35, "Einmalkosten mit Auszahlung", "=X16*X17+X18+X19", EUR0, "Dienstleister, Datenschutz, IT-Sicherheit.", False),
    (36, "Einmalkosten interne Stunden (bewertet)", "=(X20+X21)*X11", EUR0, "Schulung und interne Stunden, mit Arbeitskosten bewertet.", False),
    (37, "Einmalkosten gesamt", "=X35+X36", EUR0, "", False),
    (38, "Amortisation kalkulatorisch", '=IF(X32>0,X37/X32,"nicht erreichbar")', DEC1 + '" Monate"', "Nach wie vielen Monaten die bewertete Ersparnis alle Einmalkosten deckt.", True),
    (39, "Amortisation zahlungswirksam", '=IF(X34>0,X35/X34,"nicht erreichbar")', DEC1 + '" Monate"', "Nach wie vielen Monaten die tatsächlich gesparten Auszahlungen die Einmalauszahlungen decken.", True),
    (40, "Ergebnis nach Betrachtungszeitraum, kalkulatorisch", "=X24*X32-X37", EUR0, "", False),
    (41, "Ergebnis nach Betrachtungszeitraum, zahlungswirksam", "=X24*X34-X35", EUR0, "", False),
    (42, "Nettoeffekt je zusätzlichem Vorgang", "=(X7-X8-X9*X10)/60*X11-X12", EUR, "Ist dieser Wert negativ, lohnt sich die Lösung bei keiner Menge.", False),
    (43, "Mindestmenge, ab der sich der laufende Betrieb trägt", '=IF(X42>0,ROUNDUP(((X14+X15/12)*X11+X13)/X42,0),"nie")', INT + '" Vorgänge/Monat"', "Break-even ohne Einmalkosten.", False),
]
for r, txt, fml, fmt, note, key in ERG:
    ws.cell(r, 1, txt).font = f_bold if key else f_norm
    for j, col in enumerate("BCD"):
        c = ws.cell(r, 2 + j, fml.replace("X", col)); c.number_format = fmt; c.border = box
        c.fill = fill_key if key else fill_res; c.font = f_bold if key else f_norm
    n = ws.cell(r, 6, note); n.font = f_note; n.alignment = wrap
for r, unit in zip(range(27, 44), ["h/Monat"] * 3 + ["€/Monat"] * 5 + ["€"] * 3 + ["Monate"] * 2 + ["€"] * 2 + ["€", "Vorgänge"]):
    ws.cell(r, 5, unit).font = f_norm

ws.cell(45, 1, "Einschätzung").font = f_bold
ws.cell(46, 1, "Hinweise zu den Eingaben").font = f_bold
for col in "BCD":
    X = col
    e = ws[f"{X}45"]
    e.value = (f'=IFERROR(IF({X}6=0,"Keine Vorgänge eingetragen.",IF({X}32<=0,"Rechnet sich nicht: Die bewertete Zeitersparnis '
               f'deckt die laufenden Kosten nicht.",IF({X}40<0,"Laufend positiv, aber die Einmalkosten werden im Betrachtungszeitraum '
               f'nicht zurückverdient.",IF({X}34<=0,"Kalkulatorisch lohnend. Geld spart es nur, wenn die frei werdende Zeit '
               f'tatsächlich Ausgaben ersetzt.",IF({X}41<0,"Kalkulatorisch lohnend; die Einmalauszahlungen werden im '
               f'Betrachtungszeitraum aber nicht durch tatsächlich gesparte Ausgaben zurückverdient.","Kalkulatorisch und '
               f'zahlungswirksam positiv. Annahmen im Pilotprojekt messen."))))),'
               f'"Eingaben prüfen.")')
    h = ws[f"{X}46"]
    h.value = (f'=IF(({X}6<=0)+({X}8>={X}7)+({X}9>0.5)+({X}7>240)+({X}11<15)+({X}11>150)+({X}12>1)=0,"keine",'
               f'IF({X}6<=0,"Vorgänge fehlen. ","")&IF({X}8>={X}7,"Mit KI nicht schneller als heute. ","")'
               f'&IF({X}9>0.5,"Nacharbeit über 50 % ist sehr hoch. ","")&IF({X}7>240,"Mehr als 4 Stunden je Vorgang: '
               f'Einheit Minuten prüfen. ","")&IF(OR({X}11<15,{X}11>150),"Arbeitskosten je Stunde ungewöhnlich, Blatt '
               f'Arbeitskosten prüfen. ","")&IF({X}12>1,"KI-Kosten über 1 € je Vorgang: Angabe prüfen.",""))')
    for cc in (e, h):
        cc.alignment = wrap; cc.font = f_norm; cc.border = box; cc.fill = fill_res
    ws.conditional_formatting.add(f"{X}46", FormulaRule(formula=[f'{X}46<>"keine"'], font=Font(color="C00000", bold=True)))
ws.row_dimensions[45].height = 62; ws.row_dimensions[46].height = 62
ws["A48"] = ("Kalkulatorisch = frei werdende Arbeitszeit, mit Arbeitskosten bewertet. Zahlungswirksam = was tatsächlich an "
             "Ausgaben wegfällt. Eine Lösung kann kalkulatorisch lohnen und trotzdem kein Geld sparen, solange die frei "
             "werdende Zeit nicht anders genutzt wird. Umsatz zählt nie direkt, nur der Deckungsbeitrag (Zeile 23).")
ws["A48"].font = f_note; ws["A48"].alignment = wrap; ws.merge_cells("A48:F48"); ws.row_dimensions[48].height = 42
ws.freeze_panes = "B6"

# 3. Startblatt
st = wb.create_sheet("Start", 0)
st.sheet_view.showGridLines = False
st.column_dimensions["A"].width = 4; st.column_dimensions["B"].width = 100
zeilen = [
    ("KI wirtschaftlich einsetzen", f_title),
    ("Rechenmodell zum Buch von Maximilian Rosen · Version 1.0 · Preisstand 23.10.2026 · alle Beträge netto", f_note),
    ("", None),
    ("In drei Schritten zur eigenen Rechnung", f_h2),
    ("1.  Öffnen Sie das Blatt „Mein Vorhaben“ (Link unten).", f_norm),
    ("2.  Ersetzen Sie die blau hinterlegten Beispielwerte durch Ihre eigenen Zahlen. Graue Felder rechnet die Tabelle selbst.", f_norm),
    ("3.  Lesen Sie Zeile „Einschätzung“ und die Hinweise. Rote Hinweise bedeuten: Eingaben prüfen.", f_norm),
    ("", None),
    ("→ Zum Blatt „Mein Vorhaben“", f_link),
    ("", None),
    ("Was Sie dafür brauchen", f_h2),
    ("Anzahl der Vorgänge je Monat, gemessene Minuten je Vorgang heute, Arbeitskosten je Stunde (Lohnbuchhaltung), "
     "Angebot des Dienstleisters mit einmaligen und laufenden Kosten.", f_norm),
    ("", None),
    ("Farben", f_h2),
    ("Blau hinterlegt = Eingabe · Grau = Ergebnis · Grün = wichtigstes Ergebnis · Gelb (Buchblätter) = einflussreiche Annahme", f_norm),
    ("Formeln sind gegen versehentliches Überschreiben geschützt. Schutz aufheben: Überprüfen › Blattschutz aufheben (kein Kennwort).", f_norm),
    ("", None),
    ("Weitere Blätter", f_h2),
    ("Die übrigen Blätter enthalten alle Rechnungen des Buches (Kapitel 3 bis 11). Eine Übersicht steht im Blatt „Anleitung“.", f_norm),
    ("", None),
    ("Wichtig", f_h2),
    ("Frei werdende Arbeitszeit ist nicht automatisch eingespartes Geld. Die Tabelle weist beides getrennt aus.", f_norm),
    ("Die Ergebnisse sind Modellrechnungen mit Ihren Annahmen. Messen Sie die Werte vor einer Entscheidung im Pilotprojekt "
     "(Kapitel 11). Keine Makros, keine Datenverbindungen.", f_norm),
]
for i, (t, f) in enumerate(zeilen, start=2):
    c = st.cell(i, 2, t)
    if f: c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
st["B10"] = '=HYPERLINK("#\'Mein Vorhaben\'!A1","→ Zum Blatt „Mein Vorhaben“")'

# 4. Eingaben einheitlich kennzeichnen und Formeln schützen
OFFEN = {"Preise", "Bewertungsmatrix"}   # freie Zeilen für eigene Modelle bzw. Kandidaten
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if s.title in OFFEN and not (isinstance(c.value, str) and c.value.startswith("=")) and c.row >= 5:
                c.protection = openpyxl.styles.Protection(locked=False)
            if c.font is not None and c.font.color is not None and c.font.color.rgb in ("FF0000FF", "000000FF") \
                    and not (isinstance(c.value, str) and str(c.value).startswith("=")):
                c.protection = openpyxl.styles.Protection(locked=False)
                if c.fill is None or c.fill.fgColor is None or c.fill.fgColor.rgb in ("00000000", None):
                    c.fill = fill_in
    s.protection.sheet = True
    s.protection.formatColumns = False   # Spaltenbreite änderbar
    s.protection.formatRows = False
wb.active = 0
wb.save(dst)
print("gespeichert", dst)
