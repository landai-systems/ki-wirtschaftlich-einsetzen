"""Unabhängige Kontrollrechnung des Blatts „Mein Vorhaben“ (ohne Excel-Formeln, direkt nach Anhang A des Buches).
Aufruf: python kontrollrechnung.py <neu_berechnete.xlsx>   Liest die gespeicherten Werte und vergleicht."""
import math, sys
import openpyxl

def rechne(V, m0, m1, f, mk, S, k, Sw, Wh, NT, Eh, Sext, DS, IT, Sh, Ih, r, DB, H):
    Wt = Wh + NT / 12
    T0, T1 = V * m0 / 60, (V * m1 + V * f * mk) / 60 + Wt
    KG, LK = (T0 - T1) * S, V * k + Sw
    NK, NZ = KG - LK, r * KG + DB - LK
    EB, EI = Eh * Sext + DS + IT, (Sh + Ih) * S
    EG = EB + EI
    s = (m0 - m1 - f * mk) / 60 * S - k
    return {27: T0, 28: T1, 29: T0 - T1, 30: KG, 31: LK, 32: NK, 33: r * KG, 34: NZ, 35: EB, 36: EI, 37: EG,
            38: EG / NK if NK > 0 else "nicht erreichbar", 39: EB / NZ if NZ > 0 else "nicht erreichbar",
            40: H * NK - EG, 41: H * NZ - EB, 42: s, 43: math.ceil((Wt * S + Sw) / s) if s > 0 else "nie"}

def pruefe(pfad):
    ws = openpyxl.load_workbook(pfad, data_only=True)["Mein Vorhaben"]
    fehler = 0
    for col in "BCD":
        e = [ws[f"{col}{r}"].value for r in range(6, 25)]
        soll = rechne(*e)
        for r, v in soll.items():
            ist = ws[f"{col}{r}"].value
            ok = ist == v if isinstance(v, str) else abs(ist - v) < 1e-6 * max(1, abs(v))
            if not ok:
                fehler += 1; print("ABWEICHUNG", col, r, ist, v)
    return fehler

if __name__ == "__main__":
    n = pruefe(sys.argv[1]); print("Abweichungen:", n); sys.exit(1 if n else 0)
